# -*- coding: utf-8 -*-
"""
Suggest a gender for each employee whose profile has none, FOR HUMAN REVIEW.

    python scripts/suggest-gender.py <no_gender.tsv> <out.xlsx>

This produces a review sheet. It does not write to the database, and the
output should not be loaded without somebody who knows these people checking
it first.

★ WHY THIS IS A SUGGESTION AND NOT AN IMPORT.

A first name is not a gender. It correlates with one, unevenly, and most
unevenly exactly where this workforce is: across Nigerian, Vietnamese,
Punjabi, Chinese and Arabic naming conventions there are many names that are
common to both, or that read as one gender to an outsider and the other to
anyone who knows the language.

The field is used for one thing — `lib/leave-policy.ts` gates maternity and
paternity leave on it. So the two failure modes are not symmetrical:

    left blank  -> "this leave depends on the gender recorded on your
                    profile, which is not set. Ask HR to add it."
    set wrongly -> "Maternity Leave is not available for the gender recorded
                    on this profile."

The first is a prompt to fix something. The second tells a woman she cannot
take maternity leave, in a system that sounds certain, and she has to argue
with it. A wrong value is worse than a blank one here, which is why the
uncertain names are left out rather than guessed at a coin flip.

Confidence is about the NAME, not the person: HIGH means the name is strongly
associated with one gender in the naming tradition it comes from. It is still
an inference.
"""
import sys
from collections import Counter

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Names whose gender association is strong and consistent in the tradition
# they come from. Lowercased, first token only.
HIGH_FEMALE = {
    "sally", "zita", "germaine", "hania", "eunice", "mirela", "elizabeth",
    "catherine", "chidinma", "ruchi", "christina", "lexie", "ana", "trupti",
    "precious", "ijeoma", "kelly", "abitha", "oyindamola", "olabisi",
    "margaret", "linda", "jennifer", "awele", "aminat", "adetoke", "marinaa",
    "ivy", "estefania", "chloe", "bianca", "neena", "rosalin", "lily",
    "jignasha", "jess", "leila", "saba", "xiaohong", "ayona",
}
HIGH_MALE = {
    "mohab", "charles", "victor", "kevin", "mohamed", "mohamad", "mohammad",
    "david", "kwame", "john", "sahil", "neeraj", "pranav", "ashish",
    "unnikrishnan", "mustafizur", "apratim", "luis", "mauricio", "richard",
    "brett", "achut", "akindele",
}

# Names deliberately NOT classified, with the reason. Being able to say why
# is the difference between a judgement and a shrug.
WHY_UNCERTAIN = {
    "ntami": "Cameroonian; used for both",
    "sefunmi": "Yoruba; unisex",
    "ayotola": "Yoruba; unisex",
    "adedayo": "Yoruba; commonly male but used for both",
    "anh": "Vietnamese; unisex",
    "tran": "Vietnamese; usually a family name, unisex as a given name",
    "trang": "Vietnamese; usually female but not reliably",
    "uyen": "Vietnamese; usually female but not reliably",
    "phuong": "Vietnamese; unisex",
    "jinyu": "Chinese; unisex",
    # Demoted from HIGH on review: a confident label here would be a bluff.
    "mehar": "Punjabi/Urdu; unisex",
    "gurpreet": "Punjabi; unisex",
    "tashu": "Indian; used for both",
    "nema": "used for both across several traditions",
    "igar": "origin unclear from the name alone",
    "remil": "origin unclear from the name alone",
    "taulik": "origin unclear from the name alone",
    "zeain": "origin unclear from the name alone",
    "achim": "German male, but not reliably so in this workforce",
}


def classify(first_name):
    token = (first_name or "").strip().split()[0].lower() if first_name.strip() else ""
    token = token.strip("()")
    if token in HIGH_FEMALE:
        return "FEMALE", "HIGH", "name is strongly female in its tradition"
    if token in HIGH_MALE:
        return "MALE", "HIGH", "name is strongly male in its tradition"
    return "", "UNCERTAIN", WHY_UNCERTAIN.get(token, "not confidently gendered from the name alone")


def main(tsv_path, out_path):
    rows = []
    for line in open(tsv_path, encoding="utf-8"):
        if not line.strip():
            continue
        p = line.rstrip("\n").split("\t")
        if len(p) < 5:
            continue
        emp_id, first, last, region, email = p[0], p[1], p[2], p[3], p[4]
        suggested, conf, basis = classify(first)
        rows.append([emp_id, email, f"{first} {last}".strip(), region or "-",
                     suggested, conf, basis, ""])

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Gender review"

    headers = ["Employee ID", "Email", "Name", "Region",
               "Suggested", "Confidence", "Why", "CONFIRM (M / F / leave blank)"]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"

    amber = PatternFill("solid", fgColor="FFF4E0")
    for r in rows:
        ws.append(r)
        if r[5] == "UNCERTAIN":
            for c in range(1, len(headers) + 1):
                ws.cell(row=ws.max_row, column=c).fill = amber

    for i, w in enumerate([13, 36, 28, 15, 11, 12, 44, 26], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    wb.save(out_path)

    counts = Counter(r[5] for r in rows)
    by_sug = Counter(r[4] or "(left blank)" for r in rows)
    print(f"{len(rows)} employees with no gender recorded -> {out_path}\n")
    for k in ("HIGH", "UNCERTAIN"):
        print(f"  {counts.get(k, 0):3d}  {k}")
    print()
    for k, v in by_sug.most_common():
        print(f"  {v:3d}  suggested {k}")
    print("\n  The UNCERTAIN rows are shaded. They are left blank on purpose:")
    print("  a wrong value tells someone they cannot take parental leave,")
    print("  which is worse than a blank one that tells them to ask HR.")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(2)
    main(sys.argv[1], sys.argv[2])
